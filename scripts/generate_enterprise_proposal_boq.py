import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_full_proposal_document():
    doc = docx.Document()

    # Standard Margins (0.75 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)
        
        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("OmniSign Cloud™ | Commercial Proposal, Technical Architecture & BOQ")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8)
        hrun.font.color.rgb = RGBColor(140, 140, 140)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("Commercial Proposal • Confidential • OmniSign Cloud Network")
        frun.font.name = "Arial"
        frun.font.size = Pt(8)
        frun.font.color.rgb = RGBColor(140, 140, 140)

    # Color Palette Constants
    NAVY_HEX = "0F172A"
    ORANGE_HEX = "EA580C"
    AMBER_HEX = "D97706"
    EMERALD_HEX = "059669"
    SOFT_AMBER_BG = "FFFBEB"
    SOFT_BLUE_BG = "EFF6FF"
    SOFT_GREEN_BG = "ECFDF5"
    BORDER_HEX = "E2E8F0"
    DARK_TEXT = RGBColor(15, 23, 42)
    MUTED_TEXT = RGBColor(100, 116, 139)
    WHITE = RGBColor(255, 255, 255)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=140, bottom=140, left=160, right=160):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_callout(text_list, title="KEY HIGHLIGHT", bg_hex="FFFBEB", border_hex="F59E0B"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(7.0)
        cell = tbl.cell(0, 0)
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
        tcPr.append(borders)

        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(3)
        crun_title = cp.add_run(f"★ {title}\n")
        crun_title.bold = True
        crun_title.font.name = "Arial"
        crun_title.font.size = Pt(10)
        crun_title.font.color.rgb = RGBColor(180, 83, 9)

        for line in text_list:
            lp = cell.add_paragraph()
            lp.paragraph_format.space_before = Pt(1)
            lp.paragraph_format.space_after = Pt(2)
            lrun = lp.add_run(line)
            lrun.font.name = "Arial"
            lrun.font.size = Pt(9.5)
            lrun.font.color.rgb = RGBColor(30, 41, 59)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_h1(text, icon="📌"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"{icon} {text}")
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(14.5)
        run.font.color.rgb = RGBColor(15, 23, 42)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(234, 88, 12)

    def add_body(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3.5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            br = p.add_run(bold_prefix)
            br.bold = True
            br.font.name = "Arial"
            br.font.size = Pt(9.5)
            br.font.color.rgb = DARK_TEXT
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = DARK_TEXT

    # -------------------------------------------------------------
    # 1. EXECUTIVE COVER PAGE BANNER
    # -------------------------------------------------------------
    cover_table = doc.add_table(rows=1, cols=1)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_table.autofit = False
    cover_table.columns[0].width = Inches(7.0)
    cover_cell = cover_table.cell(0, 0)
    set_cell_background(cover_cell, NAVY_HEX)
    set_cell_margins(cover_cell, top=360, bottom=360, left=260, right=260)

    cp0 = cover_cell.paragraphs[0]
    cp0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c_badge = cp0.add_run("ENTERPRISE COMMERCIAL PROPOSAL & SYSTEM SPECIFICATION\n\n")
    c_badge.font.name = "Arial"
    c_badge.font.size = Pt(9.5)
    c_badge.bold = True
    c_badge.font.color.rgb = RGBColor(245, 158, 11)

    c_title = cp0.add_run("OmniSign Cloud™\n")
    c_title.font.name = "Arial"
    c_title.font.size = Pt(26)
    c_title.bold = True
    c_title.font.color.rgb = WHITE

    c_sub = cp0.add_run("Real-Time Multi-Screen Digital Signage & DOOH Advertising Platform\n\n")
    c_sub.font.name = "Arial"
    c_sub.font.size = Pt(12)
    c_sub.font.color.rgb = RGBColor(203, 213, 225)

    c_desc = cp0.add_run("Includes: Project Scope, Technical Architecture, Real-Time Edge Playback, Bill of Quantities (BOQ), Commercial Pricing, and Client Terms & Conditions (SLA).\n")
    c_desc.font.name = "Arial"
    c_desc.font.size = Pt(9.5)
    c_desc.font.color.rgb = RGBColor(148, 163, 184)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Document Meta Table
    meta_tbl = doc.add_table(rows=2, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_tbl.autofit = False
    meta_tbl.columns[0].width = Inches(3.5)
    meta_tbl.columns[1].width = Inches(3.5)
    
    for row in meta_tbl.rows:
        for c in row.cells:
            set_cell_background(c, "F8FAFC")
            set_cell_margins(c, 80, 80, 120, 120)

    meta_tbl.cell(0, 0).paragraphs[0].add_run("Client Organization: Commercial Enterprise / Advertising Partner").bold = True
    meta_tbl.cell(0, 1).paragraphs[0].add_run("Document Version: v2.4.0 Commercial Release").bold = True
    meta_tbl.cell(1, 0).paragraphs[0].add_run("Project Scope: Multi-Screen Real-Time Signage Network")
    meta_tbl.cell(1, 1).paragraphs[0].add_run("Delivery Model: Turnkey (Hardware + Cloud CMS + AMC)")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # 2. REAL-TIME MULTI-SCREEN PLAYBACK ARCHITECTURE
    # -------------------------------------------------------------
    add_h1("Real-Time Multi-Screen Playback Engine", "🖥️")
    add_body("OmniSign Cloud is engineered specifically for physical commercial displays (Smart TVs, LED Video Walls, Portrait Totems, Outdoor Billboards) that run in borderless 4K/Full HD Fullscreen 24/7/365.")
    
    add_callout([
        "• Instant Cloud Synchronization: Displays receive live content updates and layout changes in real time (<500ms) with support for multi-zone split-screen playback.",
        "• 24/7 Offline Playback Guarantee: Media assets are cached locally on each device, ensuring non-stop fullscreen broadcast with zero blackouts even during internet outages."
    ], title="HOW REAL-TIME FULLSCREEN OPERATION WORKS", bg_hex="EFF6FF", border_hex="2563EB")

    # -------------------------------------------------------------
    # 3. DEVELOPMENT & ENGINEERING ARCHITECTURE (DEV SIDE)
    # -------------------------------------------------------------
    add_h1("Development & Engineering Architecture (Development Side)", "⚙️")
    add_body("The platform is architected using a modern, scalable cloud-edge decoupled infrastructure:")

    dev_tbl = doc.add_table(rows=6, cols=3)
    dev_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    dev_tbl.autofit = False
    dev_tbl.columns[0].width = Inches(1.8)
    dev_tbl.columns[1].width = Inches(2.2)
    dev_tbl.columns[2].width = Inches(3.0)

    dheaders = ["Engineering Layer", "Technology Stack", "Role & Implementation"]
    for i, h in enumerate(dheaders):
        c = dev_tbl.cell(0, i)
        set_cell_background(c, "FEF3C7")
        set_cell_margins(c, 90, 90, 110, 110)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(146, 64, 14)

    ddata = [
        ("Cloud CMS Frontend", "React 19, TypeScript, Tailwind CSS, Vite", "Ultra-fast responsive admin dashboard, Canvas Studio, drag-and-drop zone properties inspector."),
        ("Edge Hardware Player", "HTML5 Hardware Accelerated Canvas, CSS Grid, Web Audio", "Zero-latency 60 FPS video renderer, local IndexedDB/LocalStorage cache, watchdog service."),
        ("Real-Time Transport", "WebSockets, REST API, Server-Sent Events (SSE)", "Instant payload deployment, remote screen commands (reboot, sync, volume, emergency takeover)."),
        ("Cloud Media Vault", "Object Storage (S3 / GCP), Global CDN Edge", "Automatic 4K transcoding, VAST 4.2 video tag parser, brand safety verification."),
        ("Security & RBAC", "JWT Token Auth, SHA-256 Signature, Role Hierarchy", "Cryptographically signed Proof of Play logs, role separation (Super Admin, Content Mgr, Operator, Viewer).")
    ]

    for row_idx, data in enumerate(ddata, start=1):
        for col_idx, val in enumerate(data):
            c = dev_tbl.cell(row_idx, col_idx)
            set_cell_background(c, "FFFFFF" if row_idx % 2 == 1 else "FFFBEB")
            set_cell_margins(c, 70, 70, 90, 90)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            r.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # 4. COMPREHENSIVE PROJECT FEATURES
    # -------------------------------------------------------------
    add_h1("Comprehensive Project Features", "🌟")
    add_body(" 6-digit PIN screen pairing, hardware telemetry (CPU, RAM, Temp °C, Storage), remote force-sync and reboot.", "1. Fleet Display Manager: ")
    add_body(" Custom split-screen zones across 16:9, 9:16 portrait, and 32:9 ultra-wide video walls.", "2. Canvas Studio: ")
    add_body(" News Marquee Ticker, Timezone World Clock, Live Weather, Interactive QR Connect, and Dining Menus.", "3. Dynamic Live Widgets: ")
    add_body(" Central repository for 4K video reels, graphics, dynamic HTML5, and programmatic VAST tags.", "4. Media Vault: ")
    add_body(" 7-day calendar matrix for morning, afternoon, and evening automated campaign rotations.", "5. Dayparting Automation: ")
    add_body(" Batch deployment across display groups with 1-click instant rollback capability.", "6. 1-Click Publishing: ")
    add_body(" Cryptographic timestamped impression logs for verified client invoicing and billing.", "7. Proof of Play (PoP): ")
    add_body(" One-touch full-network override for public emergency and safety broadcasts.", "8. Emergency Takeover: ")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # 5. BILL OF QUANTITIES (BOQ) & COMMERCIAL PRICING
    # -------------------------------------------------------------
    add_h1("Bill of Quantities (BOQ) & Commercial Pricing Table", "💰")
    add_body("Below is the comprehensive Bill of Quantities (BOQ) covering hardware, cloud software licenses, installation, and annual maintenance:")

    boq_tbl = doc.add_table(rows=8, cols=5)
    boq_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    boq_tbl.autofit = False
    boq_tbl.columns[0].width = Inches(0.6)
    boq_tbl.columns[1].width = Inches(2.6)
    boq_tbl.columns[2].width = Inches(1.1)
    boq_tbl.columns[3].width = Inches(1.3)
    boq_tbl.columns[4].width = Inches(1.4)

    bheaders = ["Item", "Description & Specifications", "Qty", "Unit Price", "Total Price (USD)"]
    for i, h in enumerate(bheaders):
        c = boq_tbl.cell(0, i)
        set_cell_background(c, "0F172A")
        set_cell_margins(c, 90, 90, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.color.rgb = WHITE

    boq_data = [
        ("1.0", "OmniSign Cloud CMS Enterprise License (Per screen / Annual subscription, includes 4K CDN & 99.9% SLA)", "10 Displays", "$180.00 / yr", "$1,800.00"),
        ("2.0", "Industrial 4K Edge Media Player Box (Android 13 / Linux, Quad-Core 2.0GHz, 4GB RAM, 32GB eMMC, Wi-Fi 6 + RJ45)", "10 Units", "$145.00", "$1,450.00"),
        ("3.0", "55\" Ultra-High Brightness Commercial Display Panel (4K UHD, 700 nits, 24/7 Run-Time rated, Slim Bezel)", "10 Units", "$750.00", "$7,500.00"),
        ("4.0", "Heavy-Duty VESA Wall Mount Brackets & High-Speed 4K HDMI 2.1 Shielded Cabling", "10 Sets", "$45.00", "$450.00"),
        ("5.0", "On-site Hardware Provisioning, Screen Mounting, Network Configuration & Initial Pairing", "1 Job", "$650.00", "$650.00"),
        ("6.0", "Custom Layout Template Design & Brand Media Vault Setup", "1 Package", "$400.00", "$400.00"),
        ("7.0", "Annual Maintenance Contract (AMC) & 24/7 SLA Priority Technical Support Hotline", "1 Year", "$950.00", "$950.00")
    ]

    for row_idx, data in enumerate(boq_data, start=1):
        for col_idx, val in enumerate(data):
            c = boq_tbl.cell(row_idx, col_idx)
            set_cell_background(c, "FFFFFF" if row_idx % 2 == 1 else "F8FAFC")
            set_cell_margins(c, 70, 70, 90, 90)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            r.font.color.rgb = DARK_TEXT

    # Grand Total Row
    gt_tbl = doc.add_table(rows=1, cols=2)
    gt_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    gt_tbl.autofit = False
    gt_tbl.columns[0].width = Inches(5.6)
    gt_tbl.columns[1].width = Inches(1.4)
    gt_c1 = gt_tbl.cell(0, 0)
    gt_c2 = gt_tbl.cell(0, 1)
    set_cell_background(gt_c1, "FEF3C7")
    set_cell_background(gt_c2, "FEF3C7")
    set_cell_margins(gt_c1, 90, 90, 100, 100)
    set_cell_margins(gt_c2, 90, 90, 100, 100)
    
    p_gt1 = gt_c1.paragraphs[0]
    p_gt1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_gt1 = p_gt1.add_run("ESTIMATED PROJECT TURNKEY TOTAL:")
    r_gt1.bold = True
    r_gt1.font.name = "Arial"
    r_gt1.font.size = Pt(9.5)
    r_gt1.font.color.rgb = RGBColor(146, 64, 14)

    p_gt2 = gt_c2.paragraphs[0]
    r_gt2 = p_gt2.add_run("$13,200.00")
    r_gt2.bold = True
    r_gt2.font.name = "Arial"
    r_gt2.font.size = Pt(10.5)
    r_gt2.font.color.rgb = RGBColor(180, 83, 9)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # 6. CLIENT TERMS & CONDITIONS (SLA & GOVERNANCE)
    # -------------------------------------------------------------
    add_h1("Commercial Terms & Conditions (For Clients)", "📜")
    
    terms = [
        ("1. Service Level Agreement (SLA)", "OmniSign Cloud guarantees 99.9% uptime for the cloud CMS infrastructure. In the event of unscheduled downtime exceeding 0.1% per calendar month, pro-rated service credits will be issued to the client."),
        ("2. Media Intellectual Property & Ownership", "The client retains 100% full ownership and intellectual property rights over all uploaded media assets, video commercials, logos, and promotional content. OmniSign will not use or distribute client assets without prior written consent."),
        ("3. Offline Playback Guarantee", "Edge hardware media players are provisioned with local storage caching. In the event of an ISP network outage, displays will continue broadcasting cached playlists without interruption."),
        ("4. Payment Schedule", "Turnkey commercial terms: 50% mobilization advance upon project confirmation, 30% upon hardware delivery and screen mounting, and 20% post successful go-live testing and handover."),
        ("5. Hardware Warranty & Replacement", "All hardware edge player boxes and commercial display panels include a 1-year comprehensive on-site replacement warranty against manufacturing defects."),
        ("6. Public Safety & Emergency Override", "The client acknowledges that authorized system operators retain the right to trigger instant emergency evacuation or public safety notices over the screen network in accordance with municipal guidelines.")
    ]

    for title, text in terms:
        add_body(text, f"• {title}: ")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # 7. CLIENT SIGN-OFF & ACCEPTANCE
    # -------------------------------------------------------------
    add_h1("Project Sign-off & Acceptance", "✍️")
    add_body("By signing below, both parties agree to the project scope, technical specifications, BOQ commercial estimate, and terms & conditions outlined in this proposal.")

    sign_tbl = doc.add_table(rows=3, cols=2)
    sign_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_tbl.autofit = False
    sign_tbl.columns[0].width = Inches(3.5)
    sign_tbl.columns[1].width = Inches(3.5)

    for row in sign_tbl.rows:
        for c in row.cells:
            set_cell_background(c, "F8FAFC")
            set_cell_margins(c, 90, 90, 110, 110)

    sign_tbl.cell(0, 0).paragraphs[0].add_run("FOR SERVICE PROVIDER (OmniSign):").bold = True
    sign_tbl.cell(0, 1).paragraphs[0].add_run("FOR CLIENT / ENTERPRISE:").bold = True

    sign_tbl.cell(1, 0).paragraphs[0].add_run("Authorized Signature: __________________\nName: OmniSign Cloud Lead\nTitle: Enterprise Solutions Architect")
    sign_tbl.cell(1, 1).paragraphs[0].add_run("Authorized Signature: __________________\nName: ______________________\nTitle: _______________________")

    sign_tbl.cell(2, 0).paragraphs[0].add_run("Date: ____ / ____ / 2026")
    sign_tbl.cell(2, 1).paragraphs[0].add_run("Date: ____ / ____ / 2026")

    # Save Document
    output_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "OmniSign_Enterprise_Commercial_Proposal_and_BOQ.docx"))
    try:
        doc.save(output_path)
        print(f"Successfully generated: {output_path}")
    except PermissionError:
        alt_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "OmniSign_Enterprise_Commercial_Proposal_and_BOQ_v2.docx"))
        doc.save(alt_path)
        print(f"File was open in Word. Successfully saved to updated copy: {alt_path}")

if __name__ == "__main__":
    create_full_proposal_document()
